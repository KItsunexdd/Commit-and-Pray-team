"""Build a provenance-aware address graph from downloaded source observations.

This is a data import/verification tool, not an address-string parser.
No house/apartment numbers are generated from counts or interpolation.
"""
import collections
import datetime as dt
import hashlib
import json
import math
import os
import pathlib
import re
import sys
import unicodedata

VENDOR = pathlib.Path(os.environ.get("TEMP", "/tmp")) / "jana-address-vendor"
if VENDOR.exists(): sys.path.insert(0, str(VENDOR))
from openpyxl import load_workbook
from shapely.geometry import Point, Polygon, shape
from shapely.strtree import STRtree
from shapely.prepared import prep

ROOT = pathlib.Path(__file__).resolve().parent
DATA = ROOT.parent
CITY = {"astana": "kato:710000000", "almaty": "kato:750000000", "karaganda": "kato:351010000"}
CITY_REL = {"astana": 3087155, "almaty": 2465058, "karaganda": 3386005}
ENTITIES = {}
ALIASES = []
CONFLICTS = []
RELATIONS = []


def read(name):
    return json.loads((ROOT / name).read_text(encoding="utf-8"))


def norm(s):
    return re.sub(r"[^\w]", "", unicodedata.normalize("NFKC", str(s or "")).casefold())


def number_key(value):
    # Slash, hyphen, letters and internal spaces may distinguish real objects.
    return unicodedata.normalize("NFKC", str(value or "")).casefold().strip()


def shortid(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:20]


def distance(a, b):
    return math.hypot((a[0] - b[0]) * math.cos(math.radians((a[1] + b[1]) / 2)) * 111320,
                      (a[1] - b[1]) * 111320)


def add_conflict(kind, ids, details):
    CONFLICTS.append({"id": "conflict:" + str(len(CONFLICTS) + 1), "kind": kind,
                      "entity_ids": ids, "details": details, "resolution": "retained_for_review"})


def names(ru=None, kk=None, generic=None):
    return {"ru": ru or None, "kk": kk or None, "source_default": generic or ru or kk or None}


def source_ref(source_id, record_id, **kwargs):
    return {"source_id": source_id, "record_id": str(record_id), **kwargs}


def base_entity(eid, kind, parent, ns, refs):
    entity = {"id": eid, "type": kind, "parent_id": parent, "names": ns,
              "source_refs": refs, "verification": "single_source_observation"}
    ENTITIES[eid] = entity
    return entity


def kato_type(code, ru):
    low = ru.casefold().strip()
    if code[2:] == "0000000": return "city" if int(code[:2]) > 70 else "province"
    if re.search(r"г\.\s*а\.?$", low) or low.startswith("г.а."): return "city_administration"
    if re.search(r"р\.\s*а\.?$", low) or low.startswith("р.а."): return "district_administration"
    if re.search(r"с\.\s*о\.?$", low) or low.startswith("с.о."): return "rural_district"
    if re.search(r"[сп]\.\s*а\.?$", low) or low.startswith(("с.а.", "п.а.")): return "settlement_administration"
    if low.startswith("г."): return "city"
    if "район" in low:
        return "urban_district" if int(code[:2]) > 70 or (int(code[2:4]) < 30 and code[4:6] != "00") else "district"
    if low.startswith("с."): return "rural_locality"
    if low.startswith("п."): return "settlement"
    return "small_settlement"


def kato_clean(ru, kind):
    if kind == "city": return re.sub(r"^г\.\s*", "", ru)
    if kind in {"rural_locality", "settlement"}: return re.sub(r"^[сп]\.\s*", "", ru)
    return ru


def import_kato():
    base_entity("country:KZ", "country", None,
                names("Республика Казахстан", "Қазақстан Республикасы"), [source_ref("bns_kato", "country:KZ")])
    rows = list(load_workbook((ROOT / "raw/stat_kato_xlsx.bin").open("rb"), data_only=True, read_only=True).active.values)
    codes = {str(r[0]) for r in rows[1:] if r[0]}
    for i, row in enumerate(rows[1:], 2):
        if not row[0]: continue
        code, ab, cd, ef, hij, k, kk, ru, nn = row
        code = str(code)
        kind = kato_type(code, ru)
        candidates = [code[:6] + "000", code[:4] + "00000", code[:2] + "0000000"]
        ij = int(code[-2:])
        small_parent = code[:7] + ("00" if 2 <= ij <= 79 else "80" if 82 <= ij <= 89 else "90" if 92 <= ij <= 99 else "")
        if kind == "small_settlement" and len(small_parent) == 9 and small_parent in codes and small_parent != code:
            candidates.insert(0, small_parent)
        parent_code = next((p for p in candidates if p != code and p in codes), None)
        e = base_entity("kato:" + code, kind, "kato:" + parent_code if parent_code else "country:KZ",
                        names(ru, kk), [source_ref("bns_kato", code, sheet="katonew1", row=i)])
        e.update(kato=code, display_name_ru=kato_clean(ru, kind),
                 kato_parts={"ab": str(ab).zfill(2), "cd": str(cd).zfill(2), "ef": str(ef).zfill(2), "hij": str(hij).zfill(3), "k": str(k)},
                 parent_link_method="kato_section_6_7_small_settlement_code" if parent_code == small_parent else "kato_longest_administrative_prefix", snapshot_date="2026-09-18",
                 active_status="not_individually_verified", classification_method="source_designation_and_kato_section_6_7")
        if kind == "rural_locality":
            e["legal_settlement_category"] = "not_individually_verified"
            if len(small_parent) == 9 and small_parent in codes and small_parent != code:
                e["parent_candidates"] = [{"entity_id": "kato:" + small_parent, "method": "kato_section_6_7_code_candidate", "status": "requires_individual_verification"}]
        if code[2:] == "0000000" and kind == "city":
            e["administrative_status"] = "capital" if code[:2] == "71" else "city_of_republican_significance"
    # A city and its administration are distinct source records. City districts
    # are related to the city without replacing the KATO administrative chain.
    for e in list(ENTITIES.values()):
        if e["type"] == "urban_district":
            code = e["kato"]
            city_id = "kato:" + (code[:2] + "0000000" if int(code[:2]) > 70 else code[:4] + "10000")
            if city_id in ENTITIES:
                e["locality_id"] = city_id
                RELATIONS.append({"type": "district_of_city", "from_id": e["id"], "to_id": city_id,
                                  "method": "kato_classification", "source_ids": ["bns_kato"]})
    return len(rows) - 1


def import_population():
    province_map = {}
    for e in ENTITIES.values():
        if e["parent_id"] != "country:KZ" or e["type"] not in {"province", "city"}: continue
        key = re.sub(r"^область\s+|\s+область$|^г\.\s*", "", e["names"]["ru"], flags=re.I)
        province_map[norm(key)] = e["id"]
    wb = load_workbook((ROOT / "raw/stat_population_september.bin").open("rb"), data_only=True, read_only=True)
    for rownum, row in enumerate(wb["1."].values, 1):
        if not 8 <= rownum <= 27: continue
        key = norm(re.sub(r"^г\.\s*", "", str(row[0])))
        if key in province_map:
            ENTITIES[province_map[key]]["population"] = {"value": row[5], "as_of": "2026-09-01", "source_id": "bns_population_september", "sheet": "1.", "row": rownum}
    # The July workbook has a city/district table, scoped by province. Match
    # exact normalised official names and type; ambiguous matches stay unresolved.
    wb = load_workbook((ROOT / "raw/stat_population_localities_july.bin").open("rb"), data_only=True, read_only=True)
    province = None
    attached = 0
    for rownum, row in enumerate(wb["3"].values, 1):
        text = str(row[0] or "").strip()
        value = row[1]
        if not isinstance(value, (int, float)) or not text: continue
        key = norm(re.sub(r"^город\s+|^[гсп]\.\s*", "", text))
        if key in province_map:
            province = province_map[key]
            continue
        if not province: continue
        wanted_type = "city" if re.match(r"^(город\s+|г\.)", text) else "village" if text.startswith("с.") else "settlement" if text.startswith("п.") else None
        if wanted_type is None: continue
        code_prefix = ENTITIES[province]["kato"][:2]
        matches = [e for e in ENTITIES.values() if (e["type"] == wanted_type or (wanted_type == "village" and e["type"] in {"rural_locality", "small_settlement"})) and e.get("kato", "").startswith(code_prefix)
                   and norm(e["display_name_ru"]) == key]
        if len(matches) == 1:
            matches[0]["population"] = {"value": int(value), "as_of": "2026-07-01", "source_id": "bns_population_cities_july", "sheet": "3", "row": rownum}
            attached += 1
    return attached


def geo_role(kind):
    return "street" if (kind or "").upper() in {"УЛИЦА", "ПРОСПЕКТ", "ПЕРЕУЛОК", "ПРОЕЗД", "БУЛЬВАР", "ПЛОЩАДЬ", "ШОССЕ", "НАБЕРЕЖНАЯ", "АЛЛЕЯ", "ДОРОГА"} else "address_area"


def import_karaganda():
    geo = read("karaganda_geonims.json")
    building_data = read("karaganda_buildings.json")
    for f in geo["features"]:
        p = f["properties"]
        eid = "kar:geonim:" + str(f["id"])
        kind = p.get("geonimtype_id/name")
        parent = "kar:geonim:" + str(p["parent_geonim_id"]) if p.get("parent_geonim_id") else CITY["karaganda"]
        e = base_entity(eid, "geonim", parent, names(p.get("name"), p.get("name_kk")), [source_ref("karaganda_geonims", f["id"])])
        e.update(locality_id=CITY["karaganda"], role=geo_role(kind), geonim_type=names(kind, p.get("geonimtype_id/name_kk")),
                 external_ids={"portal_geonim_id": str(f["id"])})
        if f.get("geometry", {}).get("bbox"): e["source_bbox"] = f["geometry"]["bbox"]
    # Names of parents and address objects are explicit source foreign keys.
    # Keep missing parents as labelled stubs rather than inventing their kind.
    for f in geo["features"] + building_data["features"]:
        p = f["properties"]
        for prefix in ["parent_geonim_id", "geonim_id", "geonim_id/parent_geonim_id"]:
            gid = p.get(prefix)
            if gid is None: continue
            eid = "kar:geonim:" + str(gid)
            if eid in ENTITIES: continue
            kind = p.get(prefix + "/geonimtype_id/name")
            e = base_entity(eid, "geonim", CITY["karaganda"], names(p.get(prefix + "/name"), p.get(prefix + "/name_kk")),
                            [source_ref("karaganda_buildings" if "building_addreg_id" in p else "karaganda_geonims", f["id"])])
            e.update(locality_id=CITY["karaganda"], role=geo_role(kind), geonim_type=names(kind), stub=True,
                     verification="foreign_key_observation", external_ids={"portal_geonim_id": str(gid)})
    for f in building_data["features"]:
        p = f["properties"]
        eid = "rka:" + str(f["id"])
        gid = "kar:geonim:" + str(p["geonim_id"]) if p.get("geonim_id") is not None else None
        count = p.get("addreg-flat[building_addreg_id]")
        e = base_entity(eid, "building", gid or CITY["karaganda"], names(), [source_ref("karaganda_buildings", f["id"])])
        e.update(locality_id=CITY["karaganda"], address_object_id=gid, house_number=str(p["number"]) if p.get("number") is not None else None,
                 corpus_block=p.get("corpusblock"), building_kind="unknown", district_ids=[],
                 external_ids={"rka": str(f["id"])}, active_status="not_verified",
                 apartment_data={"status": "not_collected", "source_premise_count": count, "enumeration_complete_in_source": False,
                                 "real_world_completeness": "unknown", "apartment_count": None, "numbered_apartments_applicable": "unknown"},
                 entrance_data={"status": "not_collected"})
        e["source_feature_kind"] = "primary_registry_address_object"
        if p.get("number") is None: e["data_quality_flags"] = ["house_number_missing_in_source"]
        bbox = f.get("geometry", {}).get("bbox")
        if bbox: e["point"] = [(bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2]
    return geo, building_data


def import_premises():
    verified = {r["rca"]: r for r in read("karaganda_premise_verification.json") if "properties" in r}
    # Earlier targeted checks are also evidence, including non-residential duplicates.
    for path in (ROOT / "raw").glob("flat_passport_*.bin"):
        raw = json.loads(path.read_text(encoding="utf-8"))
        p = raw.get("properties", {})
        key = str(p.get("id", path.stem.replace("flat_passport_", "")))
        if key not in verified: verified[key] = {"rca": key, "properties": p}
    buildings = {}
    for filename in ["karaganda_flats.json", "karaganda_flats_extended.json"]:
        for record in read(filename): buildings[record["building_rca"]] = record
    building_checks = {r["rca"]: r["properties"] for r in read("karaganda_building_verification.json") if "properties" in r}
    for rca, record in buildings.items():
        bid = "rka:" + rca
        if bid not in ENTITIES: continue
        bp = building_checks.get(rca)
        if bp:
            actual = str(bp.get("actual"))
            ENTITIES[bid]["active_status"] = "active_in_source_snapshot" if actual == "1" else "inactive_in_source_snapshot" if actual == "0" else "not_verified"
            ENTITIES[bid]["source_updated_date"] = bp.get("db_updated_date")
            ENTITIES[bid]["source_refs"].append(source_ref("karaganda_building_passports", rca))
            if bp.get("full_path"): ENTITIES[bid]["source_full_address_ru"] = bp["full_path"]
            if bp.get("full_path_kk"): ENTITIES[bid]["source_full_address_kk"] = bp["full_path_kk"]
        apt_nums = collections.defaultdict(list)
        type_counts = collections.Counter()
        inactive_apartments = 0
        active_apartments = 0
        all_checked = True
        for f in record["features"]:
            p = f["properties"]
            detail = verified.get(f["id"], {})
            vp = detail.get("properties", {})
            roomid = vp.get("roomtype_id")
            roomcode = vp.get("roomtype_code")
            fullpath = vp.get("full_path", "")
            # Both code and known source label/type evidence are retained.
            if roomid == 1 and (roomcode in {None, 300}) and (not fullpath or "Квартира" in fullpath or "квартира" in fullpath):
                kind = "apartment"
            elif roomid == 2 or roomcode == 360 or "Нежилое помещение" in fullpath:
                kind = "nonresidential_premise"
            else:
                kind = "premise_unknown"
                all_checked = False
            type_counts[kind] += 1
            num = str(p.get("number", ""))
            eid = "rka:" + str(f["id"])
            e = base_entity(eid, kind, bid, names(), [source_ref("karaganda_premises", f["id"], building_rca=rca)])
            actual = str(vp.get("actual"))
            e.update(number=num, building_id=bid, locality_id=CITY["karaganda"], entrance_id=None, floor=None,
                     external_ids={"rka": str(f["id"])},
                     active_status="active_in_source_snapshot" if actual == "1" else "inactive_in_source_snapshot" if actual == "0" else "not_verified",
                     source_room_type={"id": roomid, "code": roomcode},
                     verification="source_passport_type_checked" if vp else "type_not_verified")
            if vp:
                e["source_refs"].append(source_ref("karaganda_premise_passports", f["id"]))
                e["source_updated_date"] = vp.get("db_updated_date")
                if vp.get("full_path"): e["source_full_address_ru"] = vp["full_path"]
                if vp.get("full_path_kk"): e["source_full_address_kk"] = vp["full_path_kk"]
            if kind == "apartment" and actual == "1":
                apt_nums[num].append(eid)
                active_apartments += 1
            elif kind == "apartment" and actual == "0":
                inactive_apartments += 1
        complete = bool(record["complete"] and len(record["features"]) == record.get("total_reported"))
        ENTITIES[bid]["apartment_data"].update(
            status="enumerated_source_snapshot" if complete and all_checked else "partial_or_types_unverified",
            enumeration_complete_in_source=complete, premise_types_all_checked=all_checked,
            source_premise_count=record.get("total_reported"), source_enumerated_premise_count=len(record["features"]),
            apartment_count=active_apartments, inactive_apartment_record_count=inactive_apartments,
            all_apartment_record_count=type_counts["apartment"], unique_apartment_number_count=len(apt_nums),
            nonresidential_premise_count=type_counts["nonresidential_premise"], untyped_premise_count=type_counts["premise_unknown"],
            apartment_numbers=sorted(apt_nums, key=lambda n: (not n.isdigit(), int(n) if n.isdigit() else n)),
            numbered_apartments_applicable="yes" if apt_nums else "unknown")
        repeated = {n: ids for n, ids in apt_nums.items() if len(ids) > 1}
        if repeated:
            add_conflict("duplicate_apartment_numbers", [bid], repeated)
            ENTITIES[bid]["apartment_data"]["status"] = "conflicting_source_records"
        labelled = collections.defaultdict(list)
        for f in record["features"]: labelled[str(f["properties"].get("number"))].append("rka:" + f["id"])
        collisions = {n: ids for n, ids in labelled.items() if len(ids) > 1}
        if collisions:
            CONFLICTS.append({"id": "conflict:" + str(len(CONFLICTS) + 1), "kind": "premise_number_collisions",
                              "entity_ids": [bid], "details": collisions,
                              "resolution": "separated_by_type_and_source_lifecycle" if all_checked else "type_verification_pending"})
        if ENTITIES[bid]["active_status"] == "inactive_in_source_snapshot":
            ENTITIES[bid]["apartment_data"]["status"] = "building_inactive_in_source_snapshot"
    return len(verified)


def infer_street_kind(text):
    for needle, ru, kk in [("проспект", "проспект", "даңғыл"), ("переулок", "переулок", "тұйық көше"),
                           ("бульвар", "бульвар", None), ("улица", "улица", "көше"), ("шоссе", "шоссе", None)]:
        if needle in text.casefold(): return ru
    if "көшес" in text.casefold(): return "улица"
    if "даңғыл" in text.casefold(): return "проспект"
    return None


def stripped_street(text):
    return re.sub(r"^(улица|проспект|переулок|бульвар|шоссе|ул\.|пр\.|пр-т)\s+|\s+(улица|проспект|переулок|бульвар|көшесі|даңғылы)$", "", text, flags=re.I).strip()


def explicit_numbers_in_expression(expression):
    """Read only an explicitly reported simple range; never infer one from counts."""
    if not re.fullmatch(r"\s*\d+(?:\s*-\s*\d+)?(?:\s*[;,]\s*\d+(?:\s*-\s*\d+)?)*\s*", expression):
        return set()
    result = set()
    for part in re.split(r"[;,]", expression):
        bounds = [int(p.strip()) for p in part.split("-")]
        if len(bounds) == 1: result.add(str(bounds[0]))
        elif 0 <= bounds[0] <= bounds[1] and bounds[1] - bounds[0] <= 1000:
            result.update(str(n) for n in range(bounds[0], bounds[1] + 1))
    return result


def import_osm():
    obs = read("osm_observations.json")
    boundaries = obs["boundaries"]
    city_shapes = {}
    district_shapes = []
    for city, rid in CITY_REL.items():
        found = next((b for b in boundaries if b["osm_type"] == "relation" and b["osm_id"] == rid), None)
        if found is None: raise RuntimeError("Missing city boundary: " + city)
        city_shapes[city] = prep(shape(found["geometry"]))
        ENTITIES[CITY[city]]["osm_boundary"] = {"type": "relation", "id": rid, "source_id": "osm_geofabrik"}
    districts = [e for e in ENTITIES.values() if e["type"] == "urban_district" and e.get("locality_id") in CITY.values()]
    for boundary in boundaries:
        tags = boundary["tags"]
        if boundary["osm_id"] in CITY_REL.values(): continue
        candidates = [tags.get(k) for k in ["name", "name:ru", "name:kk"] if tags.get(k)]
        def district_key(text):
            return norm(re.sub(r"район|ауданы|имени|им\.|атындағы|атын\.", "", text, flags=re.I))
        matched = [e for e in districts if any(district_key(c) in {district_key(e["names"]["ru"] or ""), district_key(e["names"]["kk"] or "")} for c in candidates)]
        if len(matched) == 1:
            e = matched[0]
            geom = shape(boundary["geometry"])
            district_shapes.append((e["id"], e["locality_id"], prep(geom), boundary["osm_id"]))
            e["source_refs"].append(source_ref("osm_geofabrik", f"{boundary['osm_type']}/{boundary['osm_id']}"))
            e["verification"] = "name_corroborated_by_osm"
    kar_index = collections.defaultdict(set)
    for e in ENTITIES.values():
        if e["type"] != "geonim" or e.get("role") != "street": continue
        for value in e["names"].values():
            if value: kar_index[norm(value)].add(e["id"])
    street_groups = {}
    building_index = collections.defaultdict(list)
    for e in ENTITIES.values():
        if e["type"] == "building":
            obj = ENTITIES.get(e.get("address_object_id"))
            if obj:
                for val in obj["names"].values():
                    if val: building_index[(e["locality_id"], norm(stripped_street(val)), number_key(e["house_number"]))].append(e["id"])
    city_counts = collections.Counter()
    rejected_outside = 0
    osm_buildings = []
    entrance_items = []
    for item in obs["items"]:
        tags = item["tags"]
        point = Point(item["point"])
        cities = [c for c in item["candidate_cities"] if city_shapes[c].covers(point)]
        if len(cities) != 1:
            rejected_outside += 1
            continue
        city = cities[0]
        city_id = CITY[city]
        osmid = f"{item['osm_type']}/{item['osm_id']}"
        ref = source_ref("osm_geofabrik", osmid, object_timestamp=item["timestamp"])
        city_counts[city] += 1
        if item["osm_type"] == "node" and "entrance" in tags and tags["entrance"] != "no":
            entrance_items.append((item, city_id, ref))
        street_name = tags.get("addr:street") or (tags.get("name") if tags.get("highway") else None)
        geonim_id = None
        if street_name:
            matches = kar_index.get(norm(stripped_street(street_name)), set()) if city == "karaganda" else set()
            if len(matches) == 1:
                geonim_id = next(iter(matches))
                e = ENTITIES[geonim_id]
                if ref not in e["source_refs"]: e["source_refs"].append(ref)
                e["verification"] = "name_corroborated_by_osm"
            else:
                groupkey = (city_id, norm(street_name))
                if groupkey not in street_groups:
                    geonim_id = "osm:street:" + shortid("|".join(groupkey))
                    e = base_entity(geonim_id, "geonim", city_id,
                                    names(tags.get("name:ru") if tags.get("highway") else None,
                                          tags.get("name:kk") if tags.get("highway") else None, street_name), [ref])
                    e.update(locality_id=city_id, role="street", geonim_type=names(infer_street_kind(street_name)),
                             name_authority="community_map", grouping_method="same_city_exact_normalised_source_label",
                             identity_scope="source_label_group_not_certified_unique_physical_street")
                    street_groups[groupkey] = geonim_id
                else:
                    geonim_id = street_groups[groupkey]
                    ENTITIES[geonim_id]["source_refs"].append(ref)
                    if tags.get("highway"):
                        for lang in ["ru", "kk"]:
                            if tags.get("name:" + lang):
                                previous = ENTITIES[geonim_id]["names"].get(lang)
                                if previous and previous != tags["name:" + lang]:
                                    ENTITIES[geonim_id].setdefault("name_variants", {}).setdefault(lang, []).append(tags["name:" + lang])
                                else: ENTITIES[geonim_id]["names"][lang] = tags["name:" + lang]
            for key, value in tags.items():
                if key == "old_name" or key.startswith("old_name:"):
                    for old in value.split(";"):
                        alias = {"target_id": geonim_id, "old_name": old.strip(), "language": key.split(":", 1)[1] if ":" in key else None,
                                 "source_refs": [ref], "verification": "historical_name_reported_by_osm", "valid_until": None}
                        if alias not in ALIASES: ALIASES.append(alias)
        if not tags.get("addr:housenumber"):
            if tags.get("place") in {"suburb", "quarter", "neighbourhood"} and tags.get("name"):
                eid = "osm:" + osmid
                e = base_entity(eid, "geonim", city_id, names(tags.get("name:ru"), tags.get("name:kk"), tags.get("name")), [ref])
                e.update(locality_id=city_id, role="address_area", geonim_type=names(generic=tags["place"]), point=item["point"])
            continue
        if "entrance" in tags and not tags.get("building"):
            continue
        house = tags["addr:housenumber"]
        candidates = building_index.get((city_id, norm(stripped_street(street_name or "")), number_key(house)), []) if street_name else []
        candidates = list(dict.fromkeys(candidates))
        def block_compatible(candidate):
            registry_block = number_key(candidate.get("corpus_block"))
            mapped_block = number_key(tags.get("addr:block"))
            return registry_block == mapped_block
        near = [bid for bid in candidates if ENTITIES[bid].get("point") and block_compatible(ENTITIES[bid]) and distance(item["point"], ENTITIES[bid]["point"]) <= 40]
        if len(near) == 1:
            eid = near[0]
            e = ENTITIES[eid]
            e["source_refs"].append(ref)
            e["verification"] = "address_and_location_corroborated_by_osm" if any(r["source_id"] == "karaganda_buildings" for r in e["source_refs"]) else "multiple_osm_features_same_address_and_location"
        else:
            eid = "osm:" + osmid
            e = base_entity(eid, "building", geonim_id or city_id, names(tags.get("name:ru"), tags.get("name:kk"), tags.get("name")), [ref])
            e.update(locality_id=city_id, address_object_id=geonim_id, house_number=house, point=item["point"],
                     corpus_block=tags.get("addr:block"), building_kind="unknown", district_ids=[], external_ids={"osm": osmid}, active_status="mapped_in_source_snapshot",
                     apartment_data={"status": "missing_in_source", "apartment_count": None, "enumeration_complete_in_source": False,
                                     "real_world_completeness": "unknown", "numbered_apartments_applicable": "unknown"},
                     entrance_data={"status": "not_collected"})
            e["source_feature_kind"] = "mapped_building_outline" if tags.get("building") else "mapped_address_observation"
            if candidates:
                add_conflict("same_address_unmerged_locations", [eid] + candidates,
                             {"method": "exact_normalised_address", "reason": "location_not_unique_within_40m_or_block_not_confirmed"})
            if street_name:
                building_index[(city_id, norm(stripped_street(street_name)), number_key(house))].append(eid)
        if tags.get("building") in {"house", "detached", "bungalow", "cabin"}:
            e["building_kind"] = "individual_house_reported_by_osm"
            if e["apartment_data"]["status"] == "missing_in_source":
                e["apartment_data"]["status"] = "individual_house_no_unit_data"
        elif tags.get("building") == "apartments": e["building_kind"] = "apartment_building_reported_by_osm"
        for tag in ["building", "building:levels", "building:flats", "addr:flats", "addr:place", "addr:suburb", "addr:district", "addr:unit", "addr:postcode", "addr:block"]:
            if tag in tags: e.setdefault("osm_reported", {})[tag] = tags[tag]
        # Even explicit counts never create apartment identities or a maximum number.
        if tags.get("building:flats") == "0" and e["apartment_data"]["status"] in {"missing_in_source", "not_collected"}:
            e["apartment_data"].update(status="no_apartments_reported_by_osm", numbered_apartments_applicable="reported_no")
        for did, locality, geom, boundary_id in district_shapes:
            if locality == city_id and geom.covers(point):
                if did not in e["district_ids"]: e["district_ids"].append(did)
                e.setdefault("district_link_evidence", []).append({"district_id": did, "method": "point_in_osm_boundary", "osm_boundary_id": boundary_id})
        if item.get("polygon"):
            poly = Polygon(item["polygon"])
            if poly.is_valid: osm_buildings.append((eid, city_id, poly))
    # Match entrances only where the mapped point lies on/in a unique footprint.
    # Unmatched entrance observations are preserved without a guessed parent house.
    polygons = [it[2] for it in osm_buildings]
    tree = STRtree(polygons) if polygons else None
    verified_apartments = collections.defaultdict(list)
    for apartment in ENTITIES.values():
        if apartment["type"] == "apartment" and apartment["active_status"] == "active_in_source_snapshot":
            verified_apartments[apartment["building_id"]].append(apartment)
    for item, city_id, ref in entrance_items:
        p = Point(item["point"])
        indices = tree.query(p, predicate="intersects") if tree is not None else []
        buildings = list(dict.fromkeys(osm_buildings[int(i)][0] for i in indices if osm_buildings[int(i)][1] == city_id))
        bid = buildings[0] if len(buildings) == 1 else None
        eid = "osm:node/" + str(item["osm_id"])
        e = base_entity(eid, "entrance", bid or city_id, names(), [ref])
        tags = item["tags"]
        e.update(building_id=bid, locality_id=city_id, number=tags.get("ref") or tags.get("addr:door"),
                 entrance_kind=tags.get("entrance"), point=item["point"],
                 building_link_method="unique_osm_footprint_intersection" if bid else "unresolved", apartment_ids=[],
                 apartment_data_status="not_linked")
        if tags.get("addr:flats"):
            e["reported_apartment_expression"] = tags["addr:flats"]
            reported = explicit_numbers_in_expression(tags["addr:flats"])
            linked = [apt for apt in verified_apartments.get(bid, []) if apt["number"] in reported]
            if linked:
                e["apartment_ids"] = [apt["id"] for apt in linked]
                e["apartment_data_status"] = "partial_source_linked"
                e["apartment_link_method"] = "explicit_osm_range_intersected_with_existing_verified_registry_numbers"
                for apt in linked:
                    if apt["entrance_id"] is None:
                        apt["entrance_id"] = eid
                        apt["entrance_link_method"] = e["apartment_link_method"]
                    elif apt["entrance_id"] != eid:
                        add_conflict("apartment_multiple_reported_entrances", [apt["id"], apt["entrance_id"], eid], {})
        if tags.get("level"): e["reported_level"] = tags["level"]
        if bid:
            ENTITIES[bid]["entrance_data"] = {"status": "partial_osm_observations", "enumeration_complete_in_source": False}
    # The parsed country observations are much larger than the final selected
    # entities. Release them before serialising the resulting graph.
    return {"header": obs["header"]}, dict(city_counts), rejected_outside


def add_documented_aliases():
    for value, lang in [("Нур-Султан", "ru"), ("Нұр-Сұлтан", "kk")]:
        ALIASES.append({"target_id": CITY["astana"], "old_name": value, "language": lang,
                        "source_refs": [source_ref("rename_astana_2022", "U2200001017")],
                        "verification": "official_act_indexed_text", "act_date": "2022-09-17", "valid_until": None})
    ALIASES.append({"target_id": "kato:351011100", "old_name": "Октябрьский район", "language": "ru",
                    "source_refs": [source_ref("rename_karaganda_district", "V21K0026419")],
                    "verification": "official_act_indexed_text", "act_date": None, "valid_until": None})
    # Do not attach the Furmanov act to every street containing 'Nazarbayev':
    # matching the specific physical street remains a separate verification step.
    unique = {}
    for alias in ALIASES:
        key = (alias["target_id"], alias["old_name"], alias["language"])
        if key not in unique: unique[key] = alias
    ALIASES[:] = list(unique.values())
    for i, alias in enumerate(ALIASES, 1): alias["id"] = "historical:" + str(i)


def main():
    count = import_kato()
    pop_attached = import_population()
    geo, kar = import_karaganda()
    verified = import_premises()
    osm, city_counts, outside = import_osm()
    add_documented_aliases()
    counts = collections.Counter(e["type"] for e in ENTITIES.values())
    administrative_comparison = {
        "source_id": "bns_administrative_summary_july", "reference_date": "2026-07-01", "kato_date": "2026-09-18",
        "reference_counts": {"provinces": 17, "districts_including_urban": 195, "cities": 90, "urban_settlements": 29, "villages": 6060, "rural_districts": 2152},
        "imported_kato_counts": {"provinces": counts["province"], "districts_including_urban": counts["district"] + counts["urban_district"], "cities": counts["city"],
                                 "source_label_settlements_k1": sum(e["type"] == "settlement" and e.get("kato_parts", {}).get("k") == "1" for e in ENTITIES.values()),
                                 "rural_locality_source_labels_not_certified_village_count": counts["rural_locality"], "rural_district_label_candidates": counts["rural_district"]},
        "status": "partially_matching_definitions_and_dates_not_fully_reconciled",
        "action": "preserve_all_source_rows_do_not_claim_all_rows_are_current_independent_administrative_units",
    }
    add_conflict("administrative_aggregate_counts_differ", ["country:KZ"], administrative_comparison)
    by_city = {city: dict(collections.Counter(e["type"] for e in ENTITIES.values() if e.get("locality_id") == eid)) for city, eid in CITY.items()}
    apartment_coverage = {city: {
        "active_apartment_records_in_checked_snapshot": sum(e["type"] == "apartment" and e.get("locality_id") == city_id and e["active_status"] == "active_in_source_snapshot" for e in ENTITIES.values()),
        "buildings_with_enumerated_source_premises": sum(e["type"] == "building" and e.get("locality_id") == city_id and "apartment_numbers" in e["apartment_data"] for e in ENTITIES.values()),
        "coverage_status": "selected_buildings_only" if city == "karaganda" else "no_verified_apartment_enumeration_obtained",
    } for city, city_id in CITY.items()}
    sources = read("sources_catalog.json")
    result = {
        "schema_version": "1.0", "dataset_id": "janaaddress-kazakhstan-2026-10-16",
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(), "country_id": "country:KZ",
        "scope": {"administrative_units": "all_rows_of_downloaded_KATO_snapshot", "detailed_priority_cities": list(CITY.values()),
                  "detailed_city_boundary_basis": "OSM administrative relations; Karaganda is municipality",
                  "all_kazakhstan_addresses_complete": False, "apartments_invented_from_counts": False},
        "sources": sources, "entities": list(ENTITIES.values()), "relations": RELATIONS,
        "historical_names": ALIASES, "conflicts": CONFLICTS,
        "coverage": {"counts_by_type": dict(counts), "counts_by_priority_city": by_city,
                     "apartment_coverage_by_priority_city": apartment_coverage,
                     "kato_source_rows": count, "top_level_units": 20, "administrative_crosscheck": administrative_comparison,
                     "settlements_with_official_population_attached_from_july": pop_attached,
                     "karaganda_geonims": {"collected": len(geo["features"]), "source_total": geo["total_reported"], "complete": geo["complete"]},
                     "karaganda_buildings": {"collected": len(kar["features"]), "source_total": kar["total_reported"], "complete": kar["complete"], "continuation": kar["continuation"]},
                     "premise_passports_checked": verified, "osm_snapshot": osm["header"],
                     "osm_observations_inside_boundaries": city_counts, "osm_bbox_observations_outside_city_boundaries": outside,
                     "apartments_real_world_completeness": "not_certified", "apartment_to_entrance_mapping": "not_available_in_downloaded_register_lists"},
        "collection_queue": {
            "policy": "three_named_cities_first; other regions and cities by official population descending; unknown populations after known",
            "priority_cities": list(CITY.values()),
            "regions_by_population": [e["id"] for e in sorted((e for e in ENTITIES.values() if e["type"] == "province"), key=lambda e: -e.get("population", {}).get("value", -1))],
            "other_cities_by_population": [e["id"] for e in sorted((e for e in ENTITIES.values() if e["type"] == "city" and e["id"] not in CITY.values()), key=lambda e: (-e.get("population", {}).get("value", -1), e["id"]))],
            "other_settlements_with_known_population": [e["id"] for e in sorted((e for e in ENTITIES.values() if e["type"] in {"city", "rural_locality", "settlement", "small_settlement"} and e["id"] not in CITY.values() and e.get("population")), key=lambda e: (-e["population"]["value"], e["id"]))],
        },
        "interpretation": {"missing_apartment_number_means": "unknown_unless_membership_is_checked_against_a_complete_verified_source_snapshot",
                           "source_premise_count_is_not_apartment_count": True, "apartment_count_is_not_maximum_apartment_number": True,
                           "point_in_polygon_is_not_official_district_certification": True,
                           "kato_parent_links": "derived from available 2/4/6-digit administrative code prefixes; small-settlement subordination may need extra evidence"},
    }
    # Keep one entity per line: readable individual records without bloating a
    # six-figure address graph with hundreds of megabytes of indentation.
    path = DATA / "kazakhstan_addresses.json"
    with path.open("w", encoding="utf-8", newline="\n") as output:
        output.write("{\n")
        for i, (key, value) in enumerate(result.items()):
            if i: output.write(",\n")
            output.write(json.dumps(key) + ": ")
            if key == "entities":
                output.write("[\n")
                for entity_index, entity in enumerate(value):
                    if entity_index: output.write(",\n")
                    output.write(json.dumps(entity, ensure_ascii=False, separators=(",", ":")))
                output.write("\n]")
            else: output.write(json.dumps(value, ensure_ascii=False, indent=2))
        output.write("\n}\n")
    (ROOT / "build_summary.json").write_text(json.dumps(result["coverage"], ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["coverage"], ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
